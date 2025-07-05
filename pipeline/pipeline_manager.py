import datetime
import hashlib
import json
import math
import os
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

import fitz  # pymupdf
import pythoncom
from PIL import Image, ImageFilter
from google.cloud import language_v1
from google.oauth2 import service_account
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from win32com.client import Dispatch

from config import ENABLE_OCR
from pipeline.steps import sanitize, metadata, storage, publish
from utils.file_utils import generate_hash

# from handlers.pdf_handler import PDFHandler

AUDIT_LOG = "logs/audit_log.jsonl"


def run_pdf_pipeline(file_path, user_id, doc_id):
    flattened = sanitize.flatten_pdf(file_path)
    cleaned = metadata.remove_pdf_metadata(flattened)

    watermark_found = False
    detected_phrases = []

    if ENABLE_OCR:
        pattern_found, pattern_phrases = sanitize.detect_watermarks(cleaned)
        heuristic_found, heuristic_phrases = sanitize.detect_watermarks_heuristic(cleaned)

        watermark_found = pattern_found or heuristic_found
        detected_phrases = list(set(pattern_phrases + heuristic_phrases))

    print(detected_phrases)
    if watermark_found:
        cleaned = sanitize.remove_text_watermarks(cleaned, detected_phrases)

    # branded = watermark.add_visible_and_hidden_watermark(cleaned, user_id)
    branded = cleaned
    print("[STORED] FILE PATH", branded)
    stored = storage.save_file(branded, user_id, doc_id)
    publish.publish_file(stored)

    audit_entry = {
        "original_hash": generate_hash(file_path),
        "cleaned_hash": generate_hash(stored),
        "uploader_id": user_id,
        "watermark_detected": watermark_found,
        "phrases": detected_phrases,
        "timestamp": datetime.utcnow().isoformat()
    }

    os.makedirs(os.path.dirname(AUDIT_LOG), exist_ok=True)
    with open(AUDIT_LOG, "a", encoding="utf-8") as f:
        f.write(json.dumps(audit_entry) + "\n")

    return stored, audit_entry["cleaned_hash"], watermark_found


def run_word_pipeline(file_path, user_id, doc_id):
    docx_path = sanitize.run_word_sanitize(file_path)
    pdf_path = convert_to_pdf(docx_path)
    return run_pdf_pipeline(pdf_path, user_id, doc_id)


def run_excel_pipeline(file_path, user_id, doc_id):
    print(f"[EXCEL] Обробка Excel-файлу: {file_path}")

    # --- 1. Нормалізація імені ---
    base = os.path.basename(file_path)
    name, ext = os.path.splitext(base)
    ext = ext.lower()

    if name.endswith(".xlsx"):
        name = name[:-5]
    if name.endswith(".xls"):
        name = name[:-4]

    cleaned_name = f"{name}_cleaned.xlsx"
    output_dir = os.path.join("uploads", user_id, doc_id, "cleaned")
    os.makedirs(output_dir, exist_ok=True)
    cleaned_path = os.path.join(output_dir, cleaned_name)

    # --- 2. Вилучення метаданих + водяний знак ---
    if ext == ".xlsx":
        # Копіюємо файл
        shutil.copy2(file_path, cleaned_path)

        # Очистка метаданих
        _remove_xlsx_metadata(cleaned_path)

        # Додаємо водяний знак
        _add_watermark_xlsx(cleaned_path, user_id)

    elif ext == ".xls":
        if Dispatch is None:
            raise RuntimeError("❌ win32com не встановлений — неможливо обробити .xls")

        # Обробка через Excel COM
        _convert_xls_and_clean(file_path, cleaned_path)
        _add_watermark_xlsx(cleaned_path, user_id)
    else:
        raise ValueError("Непідтримуваний формат файлу")

    return cleaned_path, _make_hash(user_id), None


def run_powerpoint_pipeline(file_path, user_id, doc_id):
    print(f"[PPT] Обробка презентації: {file_path}")

    if not file_path.lower().endswith((".ppt", ".pptx")):
        raise ValueError("Непідтримуваний формат PowerPoint")

    # Отримуємо ім’я без розширення
    original_name = os.path.splitext(os.path.basename(file_path))[0]

    temp_dir = tempfile.mkdtemp()

    cleaned_pptx = os.path.join(temp_dir, f"{original_name}_cleaned.pptx")
    cleaned_pdf = os.path.join(temp_dir, f"{original_name}_cleaned.pdf")

    powerpoint = Dispatch("PowerPoint.Application")
    # powerpoint.Visible = False

    presentation = powerpoint.Presentations.Open(
        str(Path(file_path).resolve()),
        WithWindow=False
    )

    # --- Видалення нотаток ---
    for slide in presentation.Slides:
        if slide.HasNotesPage:
            try:
                for shape in slide.NotesPage.Shapes:
                    shape.Delete()
            except Exception:
                continue

    # --- Прибирання зовнішніх посилань ---
    try:
        links = presentation.LinkSources(1)  # 1 = ppLinkedOLELinks
        if links:
            for link in links:
                presentation.BreakLink(link, 1)
    except Exception:
        pass

    # --- Очищення властивостей документа ---
    props = presentation.BuiltInDocumentProperties
    for prop in ["Author", "Title", "Subject", "Comments", "Manager", "Keywords", "Last Author"]:
        try:
            props(prop).Value = ""
        except Exception:
            pass

    # --- Прибирання логотипів з master slide ---
    for master in presentation.Designs:
        try:
            for shp in master.SlideMaster.Shapes:
                if shp.Name.lower() in ["logo", "логотип", "brand", "image", "picture"]:
                    shp.Delete()
        except Exception:
            continue

    # --- Збереження очищеної копії та конвертація в PDF ---
    presentation.SaveAs(str(cleaned_pptx), 24)  # 24 = .pptx
    presentation.SaveAs(str(cleaned_pdf), 32)  # 32 = PDF
    presentation.Close()
    powerpoint.Quit()

    # --- Прогін через PDF-пайплайн ---
    final_pdf = run_pdf_pipeline(cleaned_pdf, user_id, doc_id)

    return final_pdf


def _remove_xlsx_metadata(file_path):
    temp_zip = file_path + ".temp.zip"
    with zipfile.ZipFile(file_path, 'r') as zin:
        with zipfile.ZipFile(temp_zip, 'w') as zout:
            for item in zin.infolist():
                if item.filename not in ["docProps/core.xml", "docProps/app.xml"]:
                    zout.writestr(item, zin.read(item))
    os.replace(temp_zip, file_path)


def _add_watermark_xlsx(file_path, user_id):
    timestamp = datetime.utcnow().isoformat()

    hash_str = _make_hash(user_id, timestamp)

    wb = load_workbook(file_path)
    ws = wb.active

    cell = ws["Z100"]
    cell.value = hash_str
    cell.font = Font(color="FFFFFF")
    cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    wb.save(file_path)


def _make_hash(user_id, timestamp=None):
    if timestamp is None:
        timestamp = datetime.utcnow().isoformat()
    return hashlib.sha256(f"{user_id}_{timestamp}".encode()).hexdigest()


def _convert_xls_and_clean(input_path, output_path):
    excel = Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False

    abs_input = str(Path(input_path).resolve())
    abs_output = str(Path(output_path).resolve())

    workbook = excel.Workbooks.Open(abs_input)

    # Очищення метаданих (змінюємо властивості документа)
    workbook.BuiltinDocumentProperties("Author").Value = ""
    workbook.BuiltinDocumentProperties("Keywords").Value = ""
    workbook.BuiltinDocumentProperties("Comments").Value = ""
    workbook.BuiltinDocumentProperties("Title").Value = ""
    workbook.BuiltinDocumentProperties("Subject").Value = ""
    workbook.BuiltinDocumentProperties("Last Author").Value = ""
    workbook.BuiltinDocumentProperties("Manager").Value = ""

    # Зберігаємо як .xlsx
    workbook.SaveAs(abs_output, FileFormat=51)  # 51 = xlOpenXMLWorkbook (без макросів)
    workbook.Close(False)
    excel.Quit()


def convert_to_pdf(docx_path):
    print(f"[WORD→PDF] Converting to PDF: {docx_path}")
    pdf_path = docx_path.replace(".docx", ".pdf")

    pythoncom.CoInitialize()
    word = Dispatch("Word.Application")
    word.Visible = False

    try:
        doc = word.Documents.Open(docx_path)
        doc.SaveAs(pdf_path, FileFormat=17)  # wdFormatPDF = 17
        doc.Close(False)
    finally:
        word.Quit()
        pythoncom.CoUninitialize()

    return pdf_path


# Витяг тексту з PDF через pyMuPDF
def extract_text_from_pdf(filepath):
    doc = fitz.open(filepath)
    texts = []
    for page in doc:
        texts.append(page.get_text())
    return "\n".join(texts)


# Витяг тексту з Excel через openpyxl
def extract_text_from_excel(filepath):
    wb = load_workbook(filepath, read_only=True, data_only=True)
    text = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    text.append(str(cell))
    return " ".join(text)


# Розбивка тексту на блоки по 15 000 символів
def split_text_blocks(text, block_size=15000):
    return [text[i:i + block_size] for i in range(0, len(text), block_size)]


def generate_thumbnail_webp(pdf_source, output_path="thumbnail.webp", size=(180, 240)):
    tmp_path = pdf_source

    # Відкрити PDF
    doc = fitz.open(tmp_path)
    page = doc.load_page(0)
    pix = page.get_pixmap(dpi=150)
    doc.close()

    # Перетворити на зображення
    image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    image = image.resize(size, Image.LANCZOS)
    image.save(output_path, "WEBP", quality=85)

    # Якщо файл був тимчасовий — видалити
    if pdf_source.startswith("http"):
        os.remove(tmp_path)

    return output_path


# Google Content Safety API

def moderate_text_google(text_block):
    credentials = service_account.Credentials.from_service_account_file('./credentials/credentials.json')
    client = language_v1.LanguageServiceClient(credentials=credentials)
    doc = language_v1.Document(content=text_block, type_=language_v1.Document.Type.PLAIN_TEXT)
    response = client.moderate_text(document=doc)
    return response.moderation_categories


# Автовизначення типу файлу і зчитування тексту
def extract_text_by_extension(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext in [".xls", ".xlsx"]:
        return extract_text_from_excel(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


# Головна функція пайплайну
def analyze_file_with_google(processed_file, user_id, doc_id):
    text = extract_text_by_extension(processed_file)
    text_file_path = os.path.join("uploads", user_id, doc_id, "cleaned/content.txt")

    if not text.strip():
        raise ValueError("No text extracted from file!")

    with open(text_file_path, "w", encoding="utf-8") as f:
        f.write(text)

    blocks = split_text_blocks(text, block_size=15000)

    all_results = []
    for idx, block in enumerate(blocks):
        print(f"Analyzing block {idx + 1}/{len(blocks)} ({len(block)} chars)...")
        result = moderate_text_google(block)
        all_results.append({
            "block_index": idx,
            "result": result,
            "text_snippet": block[:500],  # Для візуалізації фрагменту
        })
    return all_results


def get_preview_page_count(num_pages):
    return max(1, math.ceil(num_pages * 0.33))


def generate_crop_file(processed_file):
    doc = fitz.open(processed_file)
    num_pages = doc.page_count

    if num_pages < 3:
        doc.close()
        return None

    preview_page_count = get_preview_page_count(num_pages)

    # Створюємо новий документ
    preview_doc = fitz.open()
    for i in range(preview_page_count):
        preview_doc.insert_pdf(doc, from_page=i, to_page=i)

    # Зберігаємо прев’ю
    preview_path = os.path.splitext(processed_file)[0] + "_preview.pdf"
    preview_doc.save(preview_path)
    preview_doc.close()
    doc.close()

    return preview_path


def generate_blurred_preview_images(pdf_path, user_id, doc_id, output_dir="blurred_preview_pages"):
    output_dir = os.path.join("uploads", user_id, doc_id, "cleaned", output_dir)
    doc = fitz.open(pdf_path)
    num_pages = doc.page_count

    if num_pages < 3:
        doc.close()
        return []

    preview_page_count = get_preview_page_count(num_pages)
    start_index = preview_page_count
    end_index = min(start_index + 5, num_pages)

    os.makedirs(output_dir, exist_ok=True)
    output_files = []

    for i in range(start_index, end_index):
        page = doc.load_page(i)
        pix = page.get_pixmap(dpi=150)  # рендеримо сторінку як зображення
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

        # Накладаємо блюр
        blurred_img = img.filter(ImageFilter.GaussianBlur(radius=8))

        # Зберігаємо у форматі .webp
        filename = os.path.splitext(os.path.basename(pdf_path))[0]
        output_filename = f"blur_{i + 1}.webp"
        output_path = os.path.join(output_dir, f"{filename}_page{i + 1}.webp")
        blurred_img.save(output_path, "WEBP", quality=80)

        output_files.append([output_filename, output_path])

    doc.close()
    return output_files
